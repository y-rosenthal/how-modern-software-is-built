"""Write one .pptx per session. Run with the project venv (needs python-pptx):

    .venv/bin/python slides/render_pptx.py <spec.json>

spec.json (written by build.py):
  {"site": "https://.../", "figures": {"fig-s1-2-1": "/abs/path.png", ...},
   "out_dir": "/abs/dist/slides", "sessions": [SESSION dicts...]}
"""
import json
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

INK = RGBColor(0x1B, 0x24, 0x30)
INK2 = RGBColor(0x4A, 0x56, 0x66)
ACCENT = RGBColor(0x0E, 0x7C, 0x86)
LINE = RGBColor(0xCB, 0xD4, 0xDF)
SURFACE2 = RGBColor(0xE9, 0xEE, 0xF4)
CODEBG = RGBColor(0xED, 0xF1, 0xF5)

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)


def strip_html(s):
    s = re.sub(r"<code>(.*?)</code>", r"\1", s)
    s = re.sub(r"<[^>]+>", "", s)
    return (s.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
             .replace("&quot;", '"').replace("&#39;", "'"))


def bold_runs(paragraph, text_html, size, color=INK):
    """Add runs to a paragraph, making <b>…</b> spans bold."""
    parts = re.split(r"(<b>.*?</b>)", text_html)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("<b>")
        run = paragraph.add_run()
        run.text = strip_html(part)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = ACCENT if bold else color
        run.font.name = "Calibri"


def add_text(slide, left, top, width, height, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return tb


def add_link_button(slide, left, top, text, url):
    tb = slide.shapes.add_textbox(left, top, Inches(2.2), Inches(0.4))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r.font.name = "Calibri"
    r.hyperlink.address = url


def add_bullets(slide, left, top, width, height, bullets, size):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(size * 0.55)
        r0 = p.add_run()
        r0.text = "•  "
        r0.font.size = Pt(size)
        r0.font.color.rgb = ACCENT
        bold_runs(p, b, size)
    return tb


def add_table(slide, left, top, width, tbl, size):
    rows, cols = len(tbl["rows"]) + 1, len(tbl["head"])
    shape = slide.shapes.add_table(rows, cols, left, top, width, Inches(0.4) * rows)
    t = shape.table
    for c, h in enumerate(tbl["head"]):
        cell = t.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = SURFACE2
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size - 2)
                r.font.bold = True
                r.font.color.rgb = INK2
                r.font.name = "Calibri"
    for ri, row in enumerate(tbl["rows"], start=1):
        for c, val in enumerate(row):
            cell = t.cell(ri, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.font.color.rgb = INK
                    r.font.name = "Calibri"
    return shape


def add_code(slide, left, top, width, code, size):
    lines = code.count("\n") + 1
    height = Inches(0.32) * lines + Inches(0.3)
    box = slide.shapes.add_shape(1, left, top, width, height)  # rectangle
    box.fill.solid()
    box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = LINE
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.2)
    first = True
    for ln in code.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.name = "Consolas"
        r.font.color.rgb = INK
    return box, height


def build(spec):
    site = spec["site"].rstrip("/") + "/"
    figs = spec["figures"]
    os.makedirs(spec["out_dir"], exist_ok=True)
    for session in spec["sessions"]:
        n = session["n"]
        prs = Presentation()
        prs.slide_width, prs.slide_height = W, H
        blank = prs.slide_layouts[6]

        # title slide
        s = prs.slides.add_slide(blank)
        add_text(s, MARGIN, Inches(2.0), W - 2 * MARGIN, Inches(0.5), f"SESSION {n} · 75 MINUTES", 14, True, ACCENT)
        add_text(s, MARGIN, Inches(2.5), W - 2 * MARGIN, Inches(1.6), session["title"], 40, True, INK)
        add_text(s, MARGIN, Inches(4.1), W - 2 * MARGIN, Inches(1.0), session.get("subtitle", ""), 18, False, INK2)
        add_text(s, MARGIN, Inches(5.3), W - 2 * MARGIN, Inches(0.5), "How Modern Software Is Built (in 2026)", 14, False, INK2)
        add_link_button(s, W - MARGIN - Inches(2.2), Inches(0.4), "Open the book ↗", f"{site}#session-{n}")

        for sec in session["sections"]:
            for k, sl in enumerate(sec["slides"]):
                s = prs.slides.add_slide(blank)
                # crumb + book link
                add_text(s, MARGIN, Inches(0.35), Inches(9), Inches(0.4),
                         f"SESSION {n} · {sec['title'].upper()}", 11, True, INK2)
                add_link_button(s, W - MARGIN - Inches(2.2), Inches(0.3), "Book ↗", f"{site}#{sec['id']}")
                add_text(s, MARGIN, Inches(0.75), W - 2 * MARGIN, Inches(1.1), strip_html(sl["title"]), 30, True, INK)
                y = Inches(1.9)
                has_fig = "fig" in sl
                if "code" in sl:
                    _, h = add_code(s, MARGIN, y, W - 2 * MARGIN, sl["code"], 14)
                    y += h + Inches(0.25)
                if "table" in sl:
                    ncols = len(sl["table"]["head"])
                    size = 13 if ncols <= 3 else 11
                    shape = add_table(s, MARGIN, y, W - 2 * MARGIN, sl["table"], size)
                    y += Inches(0.42) * (len(sl["table"]["rows"]) + 1) + Inches(0.3)
                if has_fig:
                    png = figs.get("fig-" + sl["fig"])
                    if png and os.path.exists(png):
                        avail_h = (H - y - Inches(1.6)) if sl.get("bullets") else (H - y - Inches(0.6))
                        pic = s.shapes.add_picture(png, MARGIN, y, height=avail_h)
                        if pic.width > W - 2 * MARGIN:
                            ratio = (W - 2 * MARGIN) / pic.width
                            pic.width = int(pic.width * ratio)
                            pic.height = int(pic.height * ratio)
                        pic.left = int((W - pic.width) / 2)
                        y += pic.height + Inches(0.2)
                if sl.get("bullets"):
                    size = 14 if has_fig else (16 if ("table" in sl or "code" in sl) else 20)
                    remaining = H - y - (Inches(1.2) if sl.get("takeaway") else Inches(0.5))
                    add_bullets(s, MARGIN, y, W - 2 * MARGIN, max(remaining, Inches(0.8)), sl["bullets"], size)
                if sl.get("takeaway"):
                    box = s.shapes.add_shape(1, MARGIN, H - Inches(1.25), W - 2 * MARGIN, Inches(0.8))
                    box.fill.solid()
                    box.fill.fore_color.rgb = SURFACE2
                    box.line.color.rgb = ACCENT
                    tf = box.text_frame
                    tf.word_wrap = True
                    tf.margin_left = Inches(0.2)
                    p = tf.paragraphs[0]
                    bold_runs(p, sl["takeaway"], 14, INK)
                    for r in p.runs:
                        r.font.bold = True
        path = os.path.join(spec["out_dir"], f"session-{n}.pptx")
        prs.save(path)
        print("wrote", path)


if __name__ == "__main__":
    build(json.load(open(sys.argv[1])))
